from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\raine\Documents\bestatter")
OUTPUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "resources/templates/RECHNUNG.docx"
QR_PLACEHOLDER = ROOT / "resources/payment-qr-placeholder.png"
TEAL, LIGHT, SUBTLE = "164E59", "E8F1F2", "F3F7F7"
MUTED = RGBColor(82, 103, 110)


def font(run, size=10, bold=False, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size, run.bold = Pt(size), bold
    if color:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tc_mar.getparent() is None:
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            tc_mar.append(node)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_cm, alignment=WD_TABLE_ALIGNMENT.LEFT):
    table.autofit = False
    table.alignment = alignment
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW") or OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(int(Cm(w).emu / 635) for w in widths_cm)))
    if tbl_w.getparent() is None:
        tbl_pr.append(tbl_w)
    for old in list(tbl_pr.findall(qn("w:tblInd"))):
        tbl_pr.remove(old)
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(Cm(width).emu / 635)))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            dxa = int(Cm(width).emu / 635)
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(dxa))


def add_cell_text(cell, text, size=8.5, bold=False, color=None, align=None):
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    font(paragraph.add_run(text), size, bold, color)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
section.left_margin, section.right_margin = Cm(2), Cm(2)
section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(9.5)
normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(5), 1.05
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
doc.settings._element.append(update_fields)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("BESTATTERPLATTFORM | RECHNUNG"), 8, True, MUTED)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
font(p.add_run("{{case.invoice_type_label}}"), 22, True, RGBColor(22, 78, 89))

meta = doc.add_table(rows=6, cols=3)
meta.style = "Table Grid"
recipient = meta.cell(0, 0).merge(meta.cell(5, 0))
recipient.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
shade(recipient, SUBTLE)
rp = recipient.paragraphs[0]
rp.paragraph_format.space_after = Pt(6)
font(rp.add_run("Rechnung an"), 8, True, RGBColor(22, 78, 89))
rp = recipient.add_paragraph()
rp.paragraph_format.space_after = Pt(2)
font(rp.add_run("{{case.invoice_recipient_name}}"), 10, True)
rp = recipient.add_paragraph()
rp.paragraph_format.space_after = Pt(0)
font(rp.add_run("{{case.invoice_recipient_address}}"), 9)
for r, (label, value) in enumerate([
    ("Rechnungsnummer", "{{case.invoice_number}}"),
    ("Rechnungsdatum", "{{case.invoice_date}}"),
    ("Leistungszeitraum", "{{case.invoice_service_period}}"),
    ("Zahlbar bis", "{{case.invoice_due_date}}"),
    ("Fallnummer", "{{case.case_number}}"),
    ("Rechnungsart / Folge", "{{case.invoice_type_label}} | {{case.invoice_sequence}}"),
]):
    shade(meta.cell(r, 1), LIGHT)
    add_cell_text(meta.cell(r, 1), label, 7.5, True)
    add_cell_text(meta.cell(r, 2), value, 8.2)
    set_cant_split(meta.rows[r])
set_table_geometry(meta, [7.4, 3.2, 5.6])
for row in meta.rows:
    margins(row.cells[1], top=45, bottom=45)
    margins(row.cells[2], top=45, bottom=45)

p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(11), Pt(5)
font(p.add_run("Abgerechnete Leistungen und Auslagen"), 13, True, RGBColor(22, 78, 89))

items = doc.add_table(rows=4, cols=6)
items.style = "Table Grid"
for i, label in enumerate(["Pos", "Leistung", "Menge", "Einzel netto", "MwSt.", "Gesamt brutto"]):
    shade(items.cell(0, i), TEAL)
    add_cell_text(items.cell(0, i), label, 8, True, RGBColor(255, 255, 255), WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_header(items.rows[0])
heading = items.cell(1, 0).merge(items.cell(1, 5))
shade(heading, LIGHT)
add_cell_text(heading, "{{service.group_heading}}", 9, True, RGBColor(22, 78, 89))
set_cant_split(items.rows[1])
for i, value in enumerate(["{{service.position}}", "{{service.title}}", "{{service.quantity}}", "{{service.unit_price}}", "{{service.vat_rate}}", "{{service.gross}}"]):
    add_cell_text(items.cell(2, i), value, 8.2, False, None, None if i == 1 else WD_ALIGN_PARAGRAPH.RIGHT)
set_cant_split(items.rows[2])
subtotal_label = items.cell(3, 0).merge(items.cell(3, 4))
shade(subtotal_label, SUBTLE)
shade(items.cell(3, 5), SUBTLE)
add_cell_text(subtotal_label, "{{service.group_subtotal}}", 8.5, True, None, WD_ALIGN_PARAGRAPH.RIGHT)
add_cell_text(items.cell(3, 5), "{{service.group_subtotal_amount}}", 8.5, True, None, WD_ALIGN_PARAGRAPH.RIGHT)
set_cant_split(items.rows[3])
set_table_geometry(items, [0.9, 5.5, 1.4, 1.9, 1.3, 2.8], WD_TABLE_ALIGNMENT.CENTER)

p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(9), Pt(4)
font(p.add_run("Umsatzsteuerübersicht"), 10, True, RGBColor(22, 78, 89))
taxes = doc.add_table(rows=2, cols=3)
taxes.style = "Table Grid"
for i, label in enumerate(["Steuersatz", "Steuerpflichtiges Nettoentgelt", "Umsatzsteuer"]):
    shade(taxes.cell(0, i), LIGHT)
    add_cell_text(taxes.cell(0, i), label, 8.2, True, None, WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.CENTER)
for i, value in enumerate(["{{tax.rate}}", "{{tax.net}}", "{{tax.vat}}"]):
    add_cell_text(taxes.cell(1, i), value, 8.5, False, None, WD_ALIGN_PARAGRAPH.RIGHT)
set_repeat_header(taxes.rows[0])
set_cant_split(taxes.rows[1])
set_table_geometry(taxes, [2.4, 3.9, 3.2], WD_TABLE_ALIGNMENT.RIGHT)

# A minimal separator prevents Word from merging the two adjacent tables.
spacer = doc.add_paragraph()
spacer.paragraph_format.space_before = Pt(0)
spacer.paragraph_format.space_after = Pt(0)
spacer.paragraph_format.line_spacing = 0.1
spacer.paragraph_format.keep_with_next = True
font(spacer.add_run(" "), 1)

totals = doc.add_table(rows=5, cols=2)
totals.style = "Table Grid"
for r, (label, value) in enumerate([
    ("Entgelt netto", "{{order.taxable_net}}"),
    ("Umsatzsteuer", "{{order.total_vat}}"),
    ("Durchlaufende Posten", "{{order.pass_through_total}}"),
    ("Frühere Rechnungen", "{{case.invoice_prior_gross}}"),
    ("Rechnungsbetrag", "{{order.total_gross}}"),
]):
    if r == 4:
        shade(totals.cell(r, 0), TEAL)
        shade(totals.cell(r, 1), TEAL)
    color = RGBColor(255, 255, 255) if r == 4 else None
    add_cell_text(totals.cell(r, 0), label, 9, True, color, WD_ALIGN_PARAGRAPH.RIGHT)
    add_cell_text(totals.cell(r, 1), value, 9.5, True, color, WD_ALIGN_PARAGRAPH.RIGHT)
    set_cant_split(totals.rows[r])
    if r < 4:
        for cell in totals.rows[r].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
set_table_geometry(totals, [4.8, 3.2], WD_TABLE_ALIGNMENT.RIGHT)

p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(2)
font(p.add_run("Zahlungshinweis"), 10, True, RGBColor(22, 78, 89))
p = doc.add_paragraph("{{case.invoice_payment_instruction}} Verwendungszweck: {{case.invoice_payment_reference}}.")
for run in p.runs:
    font(run, 8.5)

payment = doc.add_table(rows=1, cols=2)
payment.style = "Table Grid"
left = payment.cell(0, 0)
shade(left, LIGHT)
for label, value in [
    ("Kontoinhaber", "{{case.bank_account_holder}}"), ("Bank", "{{case.bank_name}}"),
    ("IBAN", "{{case.bank_iban}}"), ("BIC", "{{case.bank_bic}}"),
]:
    p = left.add_paragraph() if left.text else left.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    font(p.add_run(label + ": "), 8, True)
    font(p.add_run(value), 8)
right = payment.cell(0, 1)
shade(right, SUBTLE)
p = right.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("{{case.invoice_payment_qr_label}}"), 7, True, RGBColor(22, 78, 89))
p = right.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(QR_PLACEHOLDER), width=Cm(1.8))
set_cant_split(payment.rows[0])
set_table_geometry(payment, [13.0, 3.2])

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("{{case.branch_name}} | {{case.branch_address}}"), 7.2, True, MUTED)
fp = footer.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("USt-IdNr. {{case.branch_vat_id}} | St.-Nr. {{case.branch_tax_number}} | {{case.branch_register}} | Geschäftsführung: {{case.branch_managing_directors}}"), 6.8, False, MUTED)
fp = footer.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("Bank: {{case.bank_name}} | IBAN {{case.bank_iban}} | BIC {{case.bank_bic}}"), 6.8, False, MUTED)

doc.save(OUTPUT)
print(OUTPUT)
