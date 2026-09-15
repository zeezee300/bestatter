from pathlib import Path
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "resources" / "templates"
BLUE = "00679E"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def base(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Aptos", Pt(10)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("{{case.branch_name}}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), RGBColor.from_string(BLUE)
    p.add_run("\n{{case.branch_address}}\nTelefon: {{case.branch_phone}} · E-Mail: {{case.branch_email}}")
    doc.add_paragraph("{{case.branch_name}} · {{case.branch_address}}")
    doc.add_paragraph("{{case.recipient_name}}\n{{case.recipient_address}}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("{{case.branch_city}}, {{case.current_date}}")
    h = doc.add_heading(title, level=1)
    h.style = doc.styles["Title"]
    h.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    h.runs[0].font.size = Pt(20)
    s = doc.add_paragraph(subtitle)
    s.runs[0].bold = True
    s.runs[0].font.color.rgb = RGBColor.from_string("52636B")
    return doc

def facts(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5), Cm(10.3)
        cells[0].vertical_alignment = cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cells[0], "EAF5F8")
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(value)
    doc.add_paragraph()

def finish(doc, note, source=""):
    doc.add_paragraph("Mit freundlichen Grüßen")
    doc.add_paragraph("{{case.branch_name}}\n{{case.responsible_employee}}")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(note)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor.from_string("607D8B")
    if source:
        p = doc.add_paragraph("Fachhinweis/Quelle: " + source)
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor.from_string("607D8B")

def save(doc, name):
    doc.core_properties.title = name.replace("_", " ").title()
    doc.core_properties.subject = "Bestatter-App – bearbeitbare Standardvorlage"
    doc.save(OUT / name)

doc = base("Abmeldung des Rundfunkbeitrags", "Mitteilung eines Sterbefalls – Vorbereitung für den Beitragsservice")
facts(doc, [("Verstorbene Person", "{{case.first_name}} {{case.last_name}}"), ("Geburtsdatum", "{{case.date_of_birth}}"), ("Sterbedatum", "{{case.date_of_death}}"), ("Letzte Anschrift", "{{case.last_residence}}"), ("Beitragsnummer", "{{case.broadcast_contribution_number}}")])
doc.add_paragraph("hiermit teilen wir den Sterbefall der oben genannten Person mit und bitten um Beendigung des zugehörigen Beitragskontos zum maßgeblichen Zeitpunkt.")
doc.add_paragraph("Als Anlage ist eine Kopie der Sterbeurkunde vorgesehen. Bitte richten Sie Rückfragen an die unten genannte Niederlassung.")
finish(doc, "Arbeitsvorlage der Bestatter-App. Vor Nutzung sind Empfänger, Beitragsnummer und Anlagen zu prüfen. Sie ersetzt nicht das offizielle Online-Formular.", "www.rundfunkbeitrag.de/buergerinnen-und-buerger/informationen/sterbefall")
save(doc, "RUNDFUNKBEITRAG_ABMELDUNG.docx")

doc = base("Mitteilung und Vertragsbeendigung", "Allgemeine Vorlage für Versorger, Telekommunikation, Vereine und sonstige Vertragspartner")
facts(doc, [("Verstorbene Person", "{{case.first_name}} {{case.last_name}}"), ("Sterbedatum", "{{case.date_of_death}}"), ("Letzte Anschrift", "{{case.last_residence}}"), ("Kunden-/Vertragsnummer", "{{case.contract_reference}}")])
doc.add_paragraph("hiermit teilen wir den Tod der oben genannten Person mit. Bitte prüfen Sie die Beendigung beziehungsweise Umstellung des Vertrags und bestätigen Sie die weiteren erforderlichen Schritte schriftlich.")
doc.add_paragraph("Bitte erstellen Sie keine Erstattung oder Auszahlung an das Bestattungsunternehmen. Berechtigte Empfänger sind gesondert festzustellen.")
finish(doc, "Universelle Arbeitsvorlage. Vertragsart, Vollmacht, Kündigungsregel und Anlagen müssen im Einzelfall geprüft werden.")
save(doc, "VERTRAGSABMELDUNG_STANDARD.docx")

doc = base("Mitteilung eines Sterbefalls an ein Kreditinstitut", "Vorbereitendes Anschreiben – keine Verfügung über Konten oder Nachlass")
facts(doc, [("Verstorbene Person", "{{case.first_name}} {{case.last_name}}"), ("Geburtsdatum", "{{case.date_of_birth}}"), ("Sterbedatum", "{{case.date_of_death}}"), ("Letzte Anschrift", "{{case.last_residence}}"), ("Kunden-/Kontoreferenz", "{{case.bank_customer_reference}}")])
doc.add_paragraph("hiermit wird der Sterbefall der oben genannten Person mitgeteilt. Bitte informieren Sie die legitimierten Angehörigen beziehungsweise die Nachlassvertretung über die benötigten Nachweise und das weitere Verfahren.")
doc.add_paragraph("Dieses Schreiben enthält weder einen Zahlungsauftrag noch eine Kontovollmacht.")
finish(doc, "Vorbereitende Mitteilung. Erbnachweis, Vollmacht, Kontoverbindung und Datenschutzberechtigung sind durch das Kreditinstitut zu prüfen.")
save(doc, "NACHLASSMITTEILUNG_BANK.docx")

doc = base("Sterbefallmitteilung an den Arbeitgeber", "Bitte um Information zu offenen Unterlagen und Ansprüchen")
facts(doc, [("Verstorbene Person", "{{case.first_name}} {{case.last_name}}"), ("Geburtsdatum", "{{case.date_of_birth}}"), ("Sterbedatum", "{{case.date_of_death}}"), ("Arbeitgeber / Personalnummer", "{{case.employer_reference}}")])
doc.add_paragraph("hiermit teilen wir den Tod der oben genannten Person mit. Bitte informieren Sie die berechtigten Angehörigen oder die Nachlassvertretung über noch erforderliche Unterlagen und gegebenenfalls offene arbeitsvertragliche Ansprüche.")
doc.add_paragraph("Bitte senden Sie personenbezogene Auskünfte ausschließlich an nachweislich berechtigte Empfänger.")
finish(doc, "Arbeitsvorlage. Beschäftigungsverhältnis, Empfängerberechtigung und erforderliche Nachweise sind vor Versand zu prüfen.")
save(doc, "ARBEITGEBER_STERBEFALLMITTEILUNG.docx")

# Upgrade the existing invoice template without replacing its established layout.
invoice_path = OUT / "RECHNUNG.docx"
invoice = Document(invoice_path)
invoice.paragraphs[0].clear()
title_run = invoice.paragraphs[0].add_run("{{case.invoice_type_label}}")
title_run.bold, title_run.font.size, title_run.font.color.rgb = True, Pt(22), RGBColor.from_string(BLUE)
if not any(row.cells[0].text == "Rechnungsart / Folge" for row in invoice.tables[0].rows):
    header_row = invoice.tables[0].add_row().cells
    header_row[0].text, header_row[1].text = "Rechnungsart / Folge", "{{case.invoice_type_label}} · {{case.invoice_sequence}}"
    shade(header_row[0], "EAF5F8")
    header_row[0].paragraphs[0].runs[0].bold = True
if not any(row.cells[0].text == "Frühere Rechnungen" for row in invoice.tables[2].rows):
    prior_row = invoice.tables[2].add_row().cells
    prior_row[0].text, prior_row[1].text = "Frühere Rechnungen", "{{case.invoice_prior_gross}}"
    shade(prior_row[0], "EAF5F8")
    prior_row[0].paragraphs[0].runs[0].bold = True
invoice.save(invoice_path)

# Keep the dynamically generated EPC QR compact enough for the A4 invoice footer.
with zipfile.ZipFile(invoice_path, "r") as source:
    entries = {name: source.read(name) for name in source.namelist()}
xml = entries["word/document.xml"].decode("utf-8")
xml = xml.replace(' QR \\q 3 \\s 50 ', ' QR \\q 3 \\s 35 ').replace(' QR \\q 3 ', ' QR \\q 3 \\s 35 ')
entries["word/document.xml"] = xml.encode("utf-8")
with zipfile.ZipFile(invoice_path, "w", zipfile.ZIP_DEFLATED) as target:
    for name, content in entries.items():
        target.writestr(name, content)
