from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resources" / "templates"
FONT = "Aptos"
ACCENT = RGBColor(46, 92, 103)
MUTED = RGBColor(90, 103, 108)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def base_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Kondolenzliste · {{case.case_number}}")
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    return doc


def add_title(doc: Document, text: str, size: float = 24) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = ACCENT


def add_center_line(doc: Document, text: str, size: float = 12, bold: bool = False, color=None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_event_table(doc: Document) -> None:
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = ("Trauerfeier", "Uhrzeit", "Ort", "Art")
    values = (
        "{{funeral_event.date}}",
        "{{funeral_event.time}}",
        "{{funeral_event.location}}",
        "{{funeral_event.category}}",
    )
    for row, label, value in zip(table.rows, labels, values):
        set_cell_width(row.cells[0], 2100)
        set_cell_width(row.cells[1], 7000)
        for cell in row.cells:
            set_cell_margins(cell, 110, 150, 110, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.color.rgb = ACCENT
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.add_run(value)


def build_deck() -> Path:
    doc = base_document()
    add_title(doc, "Kondolenzliste")
    add_center_line(doc, "In stillem Gedenken an", 12, False, MUTED)
    add_center_line(doc, "{{case.first_name}} {{case.last_name}}", 20, True)
    add_center_line(doc, "* {{case.date_of_birth}}    † {{case.date_of_death}}", 12)
    doc.add_paragraph()
    add_event_table(doc)
    doc.add_paragraph()
    add_center_line(doc, "Wir danken allen, die ihre Anteilnahme zum Ausdruck bringen.", 11, False, MUTED)
    path = OUT / "KONDOLENZLISTE_DECKBLATT.docx"
    doc.save(path)
    return path


def build_list() -> Path:
    doc = base_document()
    add_title(doc, "Kondolenzliste", 21)
    add_center_line(doc, "{{case.first_name}} {{case.last_name}}", 15, True)
    add_center_line(
        doc,
        "Trauerfeier am {{funeral_event.date}} um {{funeral_event.time}} · {{funeral_event.location}}",
        10,
        False,
        MUTED,
    )
    doc.add_paragraph()
    table = doc.add_table(rows=31, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (650, 4300, 4150)
    headers = ("Nr.", "Name", "Anschrift / Bemerkung")
    for index, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[index])
        set_cell_margins(cell, 100, 120, 100, 120)
        shade(cell, "DCE9EC")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(headers[index])
        run.bold = True
        run.font.color.rgb = ACCENT
    set_repeat_header(table.rows[0])
    for number, row in enumerate(table.rows[1:], start=1):
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell, 90, 120, 90, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(str(number))
            else:
                p.add_run(" ")
    path = OUT / "KONDOLENZLISTE.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for generated in (build_deck(), build_list()):
        print(generated)
