from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('output/Bestatter_Funktionsuebersicht_aktuell_DU_peppig.docx'); OUT.parent.mkdir(exist_ok=True)
NAVY='12304A'; BLUE='2D6FB7'; TEAL='07877A'; GOLD='D89B28'; SKY='EAF4FB'; MINT='E9F7F4'; PALE='F7F9FB'; GRID='D4DEE7'; TEXT='253746'

def shade(c,fill):
    p=c._tc.get_or_add_tcPr(); s=p.find(qn('w:shd'))
    if s is None: s=OxmlElement('w:shd'); p.append(s)
    s.set(qn('w:fill'),fill)
def margins(c):
    p=c._tc.get_or_add_tcPr(); m=p.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); p.append(m)
    for x in ('top','start','bottom','end'):
        n=OxmlElement('w:'+x); n.set(qn('w:w'),'130'); n.set(qn('w:type'),'dxa'); m.append(n)
def txt(c,s,b=False,col=TEXT,size=9.3):
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(s); r.bold=b; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(col); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
def borders(t):
    p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        n=OxmlElement('w:'+e); n.set(qn('w:val'),'single'); n.set(qn('w:sz'),'6'); n.set(qn('w:color'),GRID); b.append(n)
    p.append(b)
def table(d,heads,rows,widths=None,header= NAVY):
    t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(heads): txt(t.rows[0].cells[i],h,True,'FFFFFF',9); shade(t.rows[0].cells[i],header)
    hp=t.rows[0]._tr.get_or_add_trPr(); hp.append(OxmlElement('w:tblHeader'))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; rp=t.rows[-1]._tr.get_or_add_trPr(); rp.append(OxmlElement('w:cantSplit'))
        for i,v in enumerate(row): txt(cells[i],str(v)); shade(cells[i], PALE if ri%2 else 'FFFFFF')
    borders(t); d.add_paragraph().paragraph_format.space_after=Pt(1); return t
def heading(d,s,l=1):
    p=d.add_paragraph(style=f'Heading {l}'); p.add_run(s); return p
def label(d,s,col=TEAL):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2); r=p.add_run(s.upper()); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(col)
def benefit(d,title,body,color=SKY):
    t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(1.7); t.columns[1].width=Inches(4.9)
    txt(t.cell(0,0),title,True,NAVY,11); txt(t.cell(0,1),body,False,TEXT,9.8); shade(t.cell(0,0),color); shade(t.cell(0,1),color)
    hp=t.rows[0]._tr.get_or_add_trPr(); hp.append(OxmlElement('w:tblHeader'))
    borders(t); d.add_paragraph().paragraph_format.space_after=Pt(0)
def bullet(d,s):
    p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(s)

d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
st=d.styles; st['Normal'].font.name='Aptos'; st['Normal'].font.size=Pt(10); st['Normal'].font.color.rgb=RGBColor.from_string(TEXT); st['Normal'].paragraph_format.space_after=Pt(6); st['Normal'].paragraph_format.line_spacing=1.06
for n,sz,col in [('Title',34,NAVY),('Heading 1',21,NAVY),('Heading 2',13,BLUE)]:
    st[n].font.name='Aptos Display'; st[n].font.size=Pt(sz); st[n].font.bold=True; st[n].font.color.rgb=RGBColor.from_string(col); st[n].paragraph_format.space_before=Pt(12); st[n].paragraph_format.space_after=Pt(5)

# cover
p=d.add_paragraph(style='Title'); p.add_run('Bestatter')
p=d.add_paragraph(); r=p.add_run('Die digitale Fallakte für moderne Bestattungshäuser'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(19); r.font.color.rgb=RGBColor.from_string(BLUE)
p=d.add_paragraph(); r=p.add_run('Funktionsübersicht aktuell  |  Version 0.45.0'); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(TEAL)
d.add_paragraph('Du führst jeden Sterbefall klar, strukturiert und nachvollziehbar – von der ersten Erfassung über die Planung bis zur Rechnung. Bestatter verbindet die Informationen, Aufgaben und Dokumente, die Dein Team im Alltag braucht.')
benefit(d,'Klarheit','Alle wichtigen Informationen zu einem Fall an einem zentralen Ort.',SKY); benefit(d,'Sicherheit','Leistungen, Kosten, Dokumente und Freigaben bleiben lückenlos nachvollziehbar.',MINT); benefit(d,'Tempo','Geführte Prozesse, Vorlagen und nächste Aktionen bringen Arbeit ins Fließen.', 'FFF4DB')
p=d.add_paragraph(); p.paragraph_format.space_before=Pt(18); r=p.add_run('Für Entscheider: '); r.bold=True; r.font.color.rgb=RGBColor.from_string(GOLD); p.add_run('Eine Lösung, die operative Entlastung und kaufmännische Kontrolle zusammenbringt.')
p=d.add_paragraph(); r=p.add_run('Stand: 5. September 2026'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('667781')
d.add_page_break()

heading(d,'1 Warum Bestatter im Alltag überzeugt',1)
d.add_paragraph('Bestatter macht aus vielen einzelnen Arbeitsschritten einen klaren, verbundenen Prozess. Dein Team weiß jederzeit, was als Nächstes ansteht, welche Informationen fehlen und welche Leistungen bereits erbracht oder abgerechnet sind.')
table(d,['Dein Vorteil','So unterstützt Dich Bestatter'],[
 ['Ein zentraler Überblick','Fallakte, Aufgaben, Termine, Dokumente, Kontakte, Leistungen und Rechnungen greifen ineinander.'],
 ['Mehr Sicherheit im Prozess','Pflichtangaben, Statuswechsel, Freigaben und Änderungen werden nachvollziehbar geführt.'],
 ['Bessere Zusammenarbeit','Zuständigkeiten, Kalender, Nextcloud-Dateien und Benachrichtigungen bleiben synchron.'],
 ['Professioneller Auftritt','Konfigurierbare DOCX-/PDF-Vorlagen und strukturierte Dokumentausgaben sorgen für einheitliche Ergebnisse.'],
], [1.8,4.8], BLUE)
heading(d,'2 Dein durchgängiger Fallprozess',1)
d.add_paragraph('Vom Erstkontakt bis zum Abschluss begleitet Dich Bestatter mit klaren Phasen und passenden Kontrollpunkten.')
table(d,['Phase','Das erledigst Du'],[
 ['Erfassen','Sterbefall, Angehörige, Auftraggeber und Stammdaten strukturiert aufnehmen.'],
 ['Planen','Leistungen auswählen, KVA erstellen, Auftrag festschreiben und Aufgaben ableiten.'],
 ['Organisieren','Termine, Trauerfeiern, Mitarbeitende, Fahrzeuge, Räume und Ausstattung koordinieren.'],
 ['Durchführen','Erbrachte Leistungen, Mengen, Fremdkosten, Auslagen und Lieferantenbelege dokumentieren.'],
 ['Ausgeben','Behörden- und Abmeldedokumente sowie Trauerdruck und weitere Vorlagen erzeugen.'],
 ['Abrechnen','Prüfung durchführen, Teil- oder Schlussrechnung erzeugen und den Vorgang sauber abschließen.'],
], [1.4,5.2], TEAL)
d.add_page_break()

heading(d,'3 Die wichtigsten Funktionen im Überblick',1)
features=[
 ('Fallakte und Stammdaten','Du verwaltest Sterbedaten, Todeszeitpunkt, Wohnsitz, Standesämter, Friedhof, Bestattungsart, Auftraggeber sowie Ehe- und Partnerschaftsdaten in einer klaren Fallakte.'),
 ('Schnellerfassung und Assistent','Du erfasst Gespräche per Text oder Sprache. Erkannte Daten und Aufgaben werden übersichtlich zur Prüfung und Bestätigung vorgeschlagen.'),
 ('Aufgaben und Workflows','Du nutzt Checklisten, Fristen, Prioritäten, Wiedervorlagen und kontrollierte Folgeaktionen für einen verlässlichen Arbeitsfluss.'),
 ('Termine und Ressourcen','Du planst Termine mit Terminarten, Vor- und Nachlauf sowie passenden Mitarbeitenden, Fahrzeugen, Räumen, Kapellen und Ausstattungen.'),
 ('Leistungskatalog und Auftrag','Du pflegst Artikel, Pakete, Preise, Mengeneinheiten und Artikelgruppen. KVA und Auftrag bilden eine belastbare kaufmännische Grundlage.'),
 ('Dokumente und Abmeldungen','Du erzeugst professionelle DOCX-/PDF-Dokumente, wählst Anlagen aus und führst Abmeldungen mit Status, Fristen und Versandnachweis.'),
 ('Eingangsrechnungen','Du erfasst Lieferantenbelege, prüfst Positionen, ordnest Kosten zu und führst den Vorgang strukturiert bis zum Abschluss.'),
 ('Rechnungen und E-Rechnung','Du erzeugst Teil- und Schlussrechnungen mit Positionsauswahl, Leistungszeitraum, QR-Zahlcode und strukturierter XML-Ausgabe.'),
 ('Benachrichtigungen und Audit','Du erhältst passende Hinweise zu Fällen, Aufgaben, Terminen, Dokumenten und Rechnungen. Änderungen bleiben im Verlauf nachvollziehbar.'),
 ('Kennzahlen und CSV-Auswertungen','Du erhältst einen schnellen Überblick über Fälle, Aufgaben, Leistungen, Rechnungsvolumen, Planbelastung und dokumentierte Istzeit.'),
]
for i,(a,b) in enumerate(features): benefit(d,a,b, SKY if i%2==0 else MINT)
d.add_page_break()

heading(d,'4 Abrechnung mit einem guten Gefühl',1)
d.add_paragraph('Bestatter verbindet Auftrag, erbrachte Leistung, Kosten und Rechnung. Die Abrechnungssicherheits-Engine prüft regelbasiert und verständlich, ob alle relevanten Positionen berücksichtigt sind.')
table(d,['Die Prüfung erkennt','Damit Du'],[
 ['Leistungen ohne Rechnungsbezug','keine abrechenbare Leistung übersiehst.'],
 ['Mengenabweichungen','beauftragte, erbrachte und fakturierte Mengen sicher vergleichst.'],
 ['Doppelabrechnungen','Rechnungspositionen eindeutig und sauber zuordnest.'],
 ['Fremdkosten ohne Bezug','Lieferantenbelege nachvollziehbar in den Fall integrierst.'],
 ['Kulanz und begründete Abweichungen','geschäftliche Entscheidungen transparent dokumentierst.'],
], [2.6,4.0], GOLD)
heading(d,'5 Führung mit belastbaren Informationen',1)
d.add_paragraph('Die Auswertungen geben Dir ein klares Bild vom Betrieb. Du erkennst, wo Arbeit anfällt, wie weit Leistungen fortgeschritten sind und welches Rechnungsvolumen bereits ausgegeben wurde.')
table(d,['Kennzahl','Deine Perspektive'],[
 ['Offene und überfällige Aufgaben','Wo braucht Dein Team heute Aufmerksamkeit?'],
 ['Termine und Planbelastung','Wie verteilt sich die geplante Arbeit?'],
 ['Dokumentierte Istzeit','Welche Aktivitäten sind bereits zeitlich erfasst?'],
 ['Leistungserbringungsgrad','Wie viel der beauftragten Leistung ist umgesetzt?'],
 ['Fakturierungsgrad und Rechnungsvolumen','Welche Leistungen sind bis zur Rechnung fortgeschritten?'],
], [2.5,4.1], BLUE)
d.add_page_break()

heading(d,'6 Integration in Deine Arbeitsumgebung',1)
d.add_paragraph('Bestatter nutzt die Stärken der bestehenden Nextcloud-Umgebung und verbindet Fachlichkeit, Dateien und Zusammenarbeit in einem konsistenten Arbeitsraum.')
table(d,['Baustein','Dein Nutzen'],[
 ['Nextcloud-Dateien','Fallordner, Dokumente, Vorlagen und Ausgaben direkt am Vorgang.'],
 ['Kalender und CalDAV','Termine synchron in persönlichen Kalendern und mit sauberer Änderungshistorie.'],
 ['Kontakte und Aufgaben','Zentrale Kontakte, persönliche Zuständigkeiten und Aufgaben ohne doppelte Pflege.'],
 ['Aktivitäten und Hinweise','Relevante Änderungen erreichen die richtigen Mitarbeitenden zur passenden Zeit.'],
 ['Responsive Oberfläche','Übersicht und Bearbeitung funktionieren auch unterwegs auf Tablet und Smartphone.'],
], [1.75,4.85], TEAL)
heading(d,'7 Dein nächster Schritt',1)
d.add_paragraph('Erlebe den Nutzen an einem realistischen Musterfall: Erfasse einen Sterbefall, plane Auftrag und Trauerfeier, erzeuge Dokumente, buche Leistungen und führe die Abrechnung bis zur geprüften Rechnung. So wird in kurzer Zeit sichtbar, wie Bestatter Dein Team entlastet und Deine Prozesse stärkt.')
benefit(d,'Bestatter verbindet','Menschen, Prozesse und Zahlen in einer gemeinsamen Fallakte.', 'FFF4DB')
benefit(d,'Bestatter schafft Überblick','Jeder Schritt ist sichtbar, verständlich und nachvollziehbar.', SKY)
benefit(d,'Bestatter stärkt Entscheidungen','Du erhältst die Informationen, die Du für einen professionellen Betrieb brauchst.', MINT)

for s in d.sections:
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.add_run('Bestatter | Funktionsübersicht aktuell  ·  '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('72808A')
d.core_properties.title='Bestatter Funktionsübersicht aktuell'; d.core_properties.subject='Positive Funktionsübersicht in DU-Form'; d.core_properties.author='Bestatter'; d.save(OUT); print(OUT.resolve())
