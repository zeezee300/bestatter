"""Repair deterministic layout details in bundled DOCX templates.

The script only changes package templates in resources/templates. Runtime
provisioning still preserves every same-named customer template in Nextcloud.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "resources" / "templates"


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    for existing in list(properties.findall(qn("w:tblW"))):
        properties.remove(existing)
    total = sum(int(Cm(width).emu / 635) for width in widths_cm)
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    properties.insert(0, table_width)
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths_cm:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(Cm(width).emu / 635)))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_width)
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(int(Cm(width).emu / 635)))


def compact_cell(cell, font_size: float = 8.0) -> None:
    margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell._tc.get_or_add_tcPr().append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "55" if side in {"top", "bottom"} else "75")
        node.set(qn("w:type"), "dxa")
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size: float = 8.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)


def ensure_order_group_rows(table) -> None:
    if any("{{service.group_heading}}" in " ".join(cell.text for cell in row.cells) for row in table.rows):
        return
    header, item = table.rows[0], table.rows[1]
    heading = table.add_row()
    heading_cell = heading.cells[0].merge(heading.cells[-1])
    set_cell_text(heading_cell, "{{service.group_heading}}", bold=True, font_size=8.0)
    set_cell_shading(heading_cell, "E7EEF0")
    header._tr.addnext(heading._tr)
    subtotal = table.add_row()
    subtotal_label = subtotal.cells[0].merge(subtotal.cells[5])
    set_cell_text(subtotal_label, "{{service.group_subtotal}}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font_size=8.0)
    set_cell_text(subtotal.cells[6], "{{service.group_subtotal_amount}}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font_size=8.0)
    set_cell_shading(subtotal_label, "F2F2F2")
    set_cell_shading(subtotal.cells[6], "F2F2F2")
    item._tr.addnext(subtotal._tr)


def ensure_order_tax_table(document, service_table, summary_table) -> None:
    if any("{{tax.rate}}" in " ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows):
        return
    heading = document.add_paragraph("{{order.tax_heading}}")
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(2)
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(9)
    tax_table = document.add_table(rows=2, cols=3)
    set_table_widths(tax_table, [4.2, 6.3, 6.2])
    for cell, text in zip(tax_table.rows[0].cells, ["MwSt.-Satz", "Steuerpflichtiges Netto", "Mehrwertsteuer"]):
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=7.5)
        set_cell_shading(cell, "D9E2F3")
    for cell, text in zip(tax_table.rows[1].cells, ["{{tax.rate}}", "{{tax.net}}", "{{tax.vat}}"]):
        set_cell_text(cell, text, align=WD_ALIGN_PARAGRAPH.RIGHT, font_size=7.5)
    summary_table._tbl.addprevious(heading._p)
    summary_table._tbl.addprevious(tax_table._tbl)


def configure_order_summary(table) -> None:
    while len(table.rows) < 4:
        total_row = table.rows[-1]
        added = table.add_row()
        total_row._tr.addprevious(added._tr)
    rows = table.rows
    replacements = [
        ("{{order.taxable_net_label}}", "{{order.taxable_net}}"),
        ("{{order.total_vat_label}}", "{{order.total_vat}}"),
        ("{{order.pass_through_label}}", "{{order.pass_through_total}}"),
        ("{{order.total_gross_label}}", "{{order.total_gross}}"),
    ]
    for index, (label, value) in enumerate(replacements):
        set_cell_text(rows[index].cells[0], label, bold=index == 3, align=WD_ALIGN_PARAGRAPH.RIGHT, font_size=8.0)
        set_cell_text(rows[index].cells[1], value, bold=index == 3, align=WD_ALIGN_PARAGRAPH.RIGHT, font_size=8.0)
        if index == 3:
            set_cell_shading(rows[index].cells[0], "D9E2F3")
            set_cell_shading(rows[index].cells[1], "D9E2F3")


def separate_adjacent_tables(table) -> None:
    following = table._tbl.getnext()
    if following is not None and following.tag == qn("w:p") and not "".join(following.itertext()).strip():
        return
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "20")
    spacing.set(qn("w:lineRule"), "exact")
    properties.append(spacing)
    paragraph.append(properties)
    table._tbl.addnext(paragraph)


def repair_invoice() -> None:
    path = TEMPLATES / "RECHNUNG.docx"
    document = Document(path)
    widths = ([4.2, 12.5], [1.0, 7.0, 1.6, 2.4, 1.5, 3.2], [4.0, 3.2], [11.7, 5.0])
    for table, table_widths in zip(document.tables, widths):
        set_table_widths(table, list(table_widths))
    for cell in document.tables[1].rows[0].cells:
        compact_cell(cell, 8.0)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cell in document.tables[1].rows[1].cells:
        compact_cell(cell, 8.0)
    separate_adjacent_tables(document.tables[1])
    for paragraph in document.paragraphs:
        if paragraph.text.strip() in {"Zahlungshinweis", "{{case.invoice_payment_instruction}}"}:
            paragraph.paragraph_format.keep_with_next = True
    prevent_row_split(document.tables[3].rows[0])
    document.save(path)


def repair_order() -> None:
    path = TEMPLATES / "BESTATTUNGSAUFTRAG.docx"
    document = Document(path)
    section = document.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Beauftragte Leistungen und Kosten":
            paragraph.text = "{{order.services_heading}}"
            paragraph.runs[0].bold = True
        elif paragraph.text.startswith("Die Einzelpreise verstehen sich netto."):
            paragraph.text = "{{order.cost_note}}"
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1
    for table in document.tables[:3]:
        for row in table.rows:
            for cell in row.cells:
                compact_cell(cell, 8.0)
    service_table = document.tables[3]
    ensure_order_group_rows(service_table)
    set_table_widths(service_table, [1.0, 2.6, 4.8, 1.5, 2.4, 1.5, 2.9])
    for row in service_table.rows:
        for cell in row.cells:
            compact_cell(cell, 7.3)
    summary_table = document.tables[4]
    configure_order_summary(summary_table)
    ensure_order_tax_table(document, service_table, summary_table)
    for table in document.tables[4:]:
        for row in table.rows:
            for cell in row.cells:
                compact_cell(cell, 7.5)
    document.save(path)


if __name__ == "__main__":
    repair_invoice()
    repair_order()
    print("RECHNUNG.docx und BESTATTUNGSAUFTRAG.docx wurden layoutseitig repariert.")
