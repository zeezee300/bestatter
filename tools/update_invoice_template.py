from pathlib import Path

from docx import Document
from docx.shared import Mm
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "resources" / "templates" / "RECHNUNG.docx"
PLACEHOLDER = ROOT / "resources" / "payment-qr-placeholder.png"
EMPTY = ROOT / "resources" / "payment-qr-empty.png"


def create_images() -> None:
    size = 420
    image = Image.new("RGB", (size, size), "white")
    draw = ImageDraw.Draw(image)
    cell = 20
    for y in range(0, size, cell):
        for x in range(0, size, cell):
            if ((x // cell) * 7 + (y // cell) * 11 + (x * y // 400)) % 5 in (0, 1):
                draw.rectangle((x, y, x + cell - 1, y + cell - 1), fill="black")
    for x, y in ((20, 20), (size - 140, 20), (20, size - 140)):
        draw.rectangle((x, y, x + 119, y + 119), fill="white", outline="black", width=20)
        draw.rectangle((x + 40, y + 40, x + 79, y + 79), fill="black")
    PLACEHOLDER.parent.mkdir(parents=True, exist_ok=True)
    image.save(PLACEHOLDER, format="PNG", optimize=False)
    Image.new("RGBA", (1, 1), (255, 255, 255, 0)).save(EMPTY, format="PNG", optimize=False)


def clear_paragraph(paragraph) -> None:
    element = paragraph._element
    for child in list(element):
        if not child.tag.endswith("}pPr"):
            element.remove(child)


def paragraphs(document):
    seen = set()
    candidates = list(document.paragraphs)
    candidates += [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    for paragraph in candidates:
        marker = id(paragraph._p)
        if marker in seen:
            continue
        seen.add(marker)
        yield paragraph


def update_template() -> None:
    document = Document(TEMPLATE)
    instruction_paragraph = next(
        (paragraph for paragraph in document.paragraphs if "invoice_due_date" in paragraph.text or "invoice_payment_reference" in paragraph.text or "invoice_payment_instruction" in paragraph.text),
        None,
    )
    qr_cell = document.tables[3].cell(0, 1) if len(document.tables) > 3 else None
    qr_paragraph = qr_cell.paragraphs[0] if qr_cell is not None else None

    if qr_paragraph is None or instruction_paragraph is None:
        raise RuntimeError("QR- oder Zahlungshinweis-Absatz wurde in RECHNUNG.docx nicht gefunden.")

    clear_paragraph(instruction_paragraph)
    run = instruction_paragraph.add_run("{{case.invoice_payment_instruction}}")
    run.font.name = "Arial"

    clear_paragraph(qr_paragraph)
    for extra_paragraph in list(qr_cell.paragraphs[1:]):
        qr_cell._tc.remove(extra_paragraph._element)
    heading = qr_paragraph.add_run("SEPA-Zahlungscode")
    heading.bold = True
    heading.font.name = "Arial"
    heading.add_break()
    qr_paragraph.add_run().add_picture(str(PLACEHOLDER), width=Mm(27))

    # Repair legacy mojibake in the bundled standard template while retaining
    # the surrounding layout and styles.
    document.tables[0].cell(0, 0).paragraphs[0].text = "Rechnungsempfänger"
    document.tables[0].cell(5, 1).paragraphs[0].text = "{{case.invoice_type_label}} · {{case.invoice_sequence}}"
    document.tables[2].cell(3, 0).paragraphs[0].text = "Frühere Rechnungen"

    document.save(TEMPLATE)


if __name__ == "__main__":
    create_images()
    update_template()
    print(TEMPLATE)
